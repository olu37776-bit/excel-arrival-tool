from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from tests.test_pipeline import _run, _write_sources


class OutputSimplificationTest(unittest.TestCase):
    def test_output_has_no_system_pivots_and_no_forced_full_calculation(self):
        with TemporaryDirectory() as temporary:
            directory = Path(temporary)
            output = directory / "result.xlsx"
            _run(_write_sources(directory, "simple", variant="first"), output)

            workbook = load_workbook(output)
            try:
                self.assertEqual([
                    "基表",
                    "RPD跨月变化",
                    "CPD跨月变化",
                    "供应需要提拉诉求清单粗表",
                    "异常清单",
                    "_tool_meta",
                ], workbook.sheetnames)
                self.assertNotIn("_summary_source", workbook.sheetnames)
                self.assertFalse(any(sheet._pivots for sheet in workbook.worksheets))
                self.assertEqual("auto", workbook.calculation.calcMode)
                self.assertFalse(workbook.calculation.forceFullCalc)
                self.assertFalse(workbook.calculation.fullCalcOnLoad)
                self.assertFalse(workbook.calculation.calcOnSave)
                base = workbook["基表"]
                formula_cells = [cell for row in base.iter_rows(min_row=2) for cell in row if cell.data_type == "f"]
                self.assertEqual(4 * (base.max_row - 1), len(formula_cells))
            finally:
                workbook.close()

            with ZipFile(output) as archive:
                names = archive.namelist()
                self.assertFalse(any(name.startswith("xl/pivotTables/") for name in names))
                self.assertFalse(any(name.startswith("xl/pivotCache/") for name in names))


if __name__ == "__main__":
    unittest.main()
