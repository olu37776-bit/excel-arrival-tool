from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import os
import shutil
import subprocess
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from revenue_tool.adapters.excel_reader import ExcelInputAdapter
from revenue_tool.adapters.excel_writer import ExcelOutputAdapter
from revenue_tool.config import load_config
from revenue_tool.domain.models import (
    BaseRow, CONTRACT_ONLY_NO_DEMAND, IssueLog, PreviousData,
)
from revenue_tool.services.comparison import compare_revenue_months, build_supply_pull_rows
from revenue_tool.services.final_revenue import INVALID_GAP_HINT
from tests.test_pipeline import CONFIG, _base_rows, _run, _write_sources
from tests.test_manual_revenue_forecast import _set_manual_inputs


class RevenueSpaceTest(unittest.TestCase):
    def test_gap_positions_formula_caches_and_invalid_amounts(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'gap.xlsx'
            cases = [(None, 0), (80, 70), (0, 150), (175, -25),
                     (-20, 170), ('80.125', 69.87), ('abc', INVALID_GAP_HINT)]
            rows = [BaseRow(dict(contract_no=f'C{i}', revenue_forecast=150,
                                 manual_revenue_month=manual))
                    for i, (manual, _) in enumerate(cases)]
            ExcelOutputAdapter().write(path, rows, [], [], [], IssueLog(), load_config(CONFIG))
            live, cached = load_workbook(path), load_workbook(path, data_only=True)
            try:
                sheet = live['基表']
                self.assertEqual(('收入空间', '最终收入预测', '收入空间差额', '调整备注'),
                                 tuple(sheet[ref].value for ref in ('D1', 'AM1', 'AN1', 'AO1')))
                self.assertEqual('A1:AO8', sheet.auto_filter.ref)
                for number, (_, expected) in enumerate(cases, 2):
                    self.assertEqual(expected, cached['基表'][f'AN{number}'].value)
                    self.assertIn(f'ROUND(D{number}-AM{number},2)', sheet[f'AN{number}'].value)
                    self.assertEqual('f', sheet[f'AN{number}'].data_type)
                    self.assertEqual('#,##0.00', sheet[f'AN{number}'].number_format)
            finally:
                live.close(); cached.close()

    def test_all_sheets_have_explicit_fixed_heights_and_remain_editable(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); path = root / 'result.xlsx'
            _run(_write_sources(root, 'src', variant='first'), path)
            workbook = load_workbook(path)
            try:
                for sheet in workbook:
                    self.assertFalse(sheet.protection.sheet)
                    self.assertEqual(22, sheet.sheet_format.defaultRowHeight)
                    self.assertTrue(sheet.sheet_format.customHeight)
                    for number in range(1, sheet.max_row + 1):
                        expected = 32 if number == 1 and sheet.sheet_state == 'visible' else 22
                        self.assertEqual(expected, sheet.row_dimensions[number].height)
                        self.assertTrue(sheet.row_dimensions[number].customHeight)
                workbook['基表']['AO2'] = '长备注测试' * 100 + '\n第二行\n第三行'
                workbook.save(path)
            finally:
                workbook.close()
            with ZipFile(path) as archive:
                for name in archive.namelist():
                    if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                        tree = ET.fromstring(archive.read(name))
                        ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        self.assertIsNone(tree.find('x:sheetProtection', ns))
                        for row in tree.findall('x:sheetData/x:row', ns):
                            self.assertEqual('1', row.get('customHeight'), (name, row.attrib))
                            self.assertIn(row.get('ht'), ('22', '32'))

    def test_auxiliary_outputs_compute_space_from_each_rows_displayed_amounts(self):
        def row(contract, legacy, monthly, month='2026-09', no_demand=False):
            values = dict(contract_no=contract, supply_center=None if no_demand else '深供',
                          legacy_amount=Decimal(legacy), monthly_new_order=Decimal(monthly),
                          revenue_forecast=999999,
                          revenue_month_rpd=None if no_demand else month,
                          revenue_month_cpd=None if no_demand else '2026-11')
            return BaseRow(values, row_kind=CONTRACT_ONLY_NO_DEMAND) if no_demand else BaseRow(values)
        previous = [row('cancel', '100', '20'), row('no-demand', '90', '2'),
                    row('resume', '0', '0', no_demand=True), row('change', '1', '2')]
        current = [row('no-demand', '8', '-3', no_demand=True),
                   row('resume', '50', '5'), row('change', '30', '40', '2026-10')]
        old = PreviousData({(r.values['contract_no'], (r.values['supply_center'] or '').casefold()): r for r in previous})
        issues = IssueLog()
        rpd = compare_revenue_months(current, old, 'rpd', 'now.xlsx', issues)
        cpd = compare_revenue_months(current, old, 'cpd', 'now.xlsx', issues)
        supply = build_supply_pull_rows(current, 'now.xlsx', issues)
        expected = {'cancel': 120, 'no-demand': 5, 'resume': 55, 'change': 70}
        for rows in (rpd, cpd, supply):
            self.assertTrue(rows)
            for record in rows:
                self.assertEqual(expected[record.values['contract_no']], record.values['revenue_forecast'])
                self.assertEqual(record.values['legacy_amount'] + record.values['monthly_new_order'],
                                 record.values['revenue_forecast'])
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'aux.xlsx'
            ExcelOutputAdapter().write(path, [], rpd, cpd, supply, issues, load_config(CONFIG))
            workbook = load_workbook(path, data_only=True)
            try:
                for name in ('RPD跨月变化', 'CPD跨月变化', '供应需要提拉诉求清单粗表'):
                    sheet = workbook[name]
                    self.assertEqual('收入空间', sheet['D1'].value)
                    for cells in sheet.iter_rows(min_row=2):
                        self.assertEqual(cells[1].value + cells[2].value, cells[3].value)
                        self.assertEqual('#,##0.00', cells[3].number_format)
            finally:
                workbook.close()

    def test_old_40_column_results_keep_manual_values_and_notes(self):
        for metadata in (False, True):
            with self.subTest(metadata=metadata), TemporaryDirectory() as tmp:
                root = Path(tmp); old, new = root / 'old.xlsx', root / 'new.xlsx'
                sources = _write_sources(root, 'src', variant='first')
                _run(sources, old)
                _set_manual_inputs(old, 'C001', 'SC-A', amount=80, rpd='2026-9', cpd='2026-10', segment_flag=False)
                workbook = load_workbook(old)
                base = workbook['基表']
                base['D1'] = '收入预测'
                for cells in base.iter_rows(min_row=2):
                    if cells[0].value == 'C001' and cells[11].value == 'SC-A':
                        base.cell(cells[0].row, 41, '上期人工备注\n仍需核对')
                        base.cell(cells[0].row, 4, 999999)
                base.delete_cols(40)  # v0.11.4: AM final, AN note, no gap.
                if metadata:
                    meta = workbook['_tool_meta']
                    for cells in list(meta.iter_rows()):
                        if cells[0].value == 'revenue_forecast':
                            cells[1].value = '收入预测'
                        if cells[0].value == 'revenue_space_gap':
                            meta.delete_rows(cells[0].row)
                else:
                    del workbook['_tool_meta']
                workbook.save(old); workbook.close()
                before = old.read_bytes()
                issues = IssueLog()
                imported = ExcelInputAdapter().read_previous(old, load_config(CONFIG), issues)
                self.assertTrue(imported.usable)
                self.assertFalse(any(i.field == 'revenue_space_gap' for i in issues.items))
                _run(sources, new, previous=old)
                workbook = load_workbook(new, data_only=True)
                try:
                    actual = _base_rows(workbook['基表'])[('C001', 'SC-A')]
                    self.assertEqual(150, actual['收入空间'])
                    self.assertEqual(80, actual['最终收入预测'])
                    self.assertEqual(70, actual['收入空间差额'])
                    self.assertEqual('上期人工备注 仍需核对', actual['调整备注'])
                    self.assertEqual('2026-09', actual['最终收入年月（按RPD）'])
                    self.assertIs(False, actual['最终收入分段类别'])
                finally:
                    workbook.close()
                self.assertEqual(before, old.read_bytes())

    def test_stale_gap_is_not_an_inherited_input(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); old, new = root / 'old.xlsx', root / 'new.xlsx'
            sources = _write_sources(root, 'src', variant='first')
            _run(sources, old)
            _set_manual_inputs(old, 'C001', 'SC-A', amount=0)
            workbook = load_workbook(old)
            for number in range(2, workbook['基表'].max_row + 1):
                workbook['基表'].cell(number, 40, '错误旧差额')
            workbook.save(old); workbook.close()
            imported = ExcelInputAdapter().read_previous(old, load_config(CONFIG), IssueLog())
            self.assertTrue(all('revenue_space_gap' not in row.values for row in imported.rows.values()))
            _run(sources, new, previous=old)
            workbook = load_workbook(new, data_only=True)
            try:
                actual = _base_rows(workbook['基表'])[('C001', 'SC-A')]
                self.assertEqual(150, actual['收入空间差额'])
            finally:
                workbook.close()

    def test_live_edit_clear_and_multiline_note_with_independent_engine(self):
        office = shutil.which('libreoffice') or shutil.which('soffice')
        python = Path('/usr/bin/python3')
        uno_available = python.exists() and subprocess.run(
            [str(python), '-c', 'import uno'], capture_output=True,
        ).returncode == 0
        if not office or not uno_available:
            if os.environ.get('REQUIRE_FORMULA_ENGINE') == '1':
                self.fail('Live-edit verification requires LibreOffice and python3-uno')
            self.skipTest('Live-edit check runs in the Linux verification job')
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'live.xlsx'
            values = dict(contract_no='LIVE', legacy_amount=100, monthly_new_order=50,
                          revenue_forecast=150, adjustment_note='初始备注')
            ExcelOutputAdapter().write(path, [BaseRow(values)], [], [], [], IssueLog(), load_config(CONFIG))
            helper = Path(__file__).parent / 'support' / 'verify_live_revenue_space.py'
            completed = subprocess.run([str(python), str(helper), office, str(path)],
                                       capture_output=True, text=True, timeout=90)
            self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
