"""Issue #50: empty sheets must not abort read-only XLSX import."""
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from revenue_tool.adapters.excel_reader import ExcelInputAdapter, _open_workbook
from revenue_tool.config import load_config
from revenue_tool.domain.models import IssueLog, SourceFiles, WorkbookReadError
from tests.test_grouped_pivot_input import NS, add_discrete_group, rewrite_package
from tests.test_manual_revenue_forecast import _set_manual_inputs
from tests.test_pipeline import CONFIG, _base_rows, _run, _write_sources
from tests.test_user_pivot_regions import add_user_pivots


def add_empty_sheets(path):
    workbook = load_workbook(path)
    try:
        for state in ('visible', 'hidden', 'veryHidden'):
            workbook.create_sheet('空白-' + state).sheet_state = state
        workbook.save(path)
    finally:
        workbook.close()


def alter_dimensions(path, dimension, empty_rows=False):
    def transform(name, data):
        if name != 'xl/worksheets/sheet1.xml':
            return data
        root = ET.fromstring(data)
        element = root.find(f'{{{NS}}}dimension')
        if dimension is None:
            root.remove(element)
        else:
            element.set('ref', dimension)
        if empty_rows:
            sheet_data = root.find(f'{{{NS}}}sheetData')
            ET.SubElement(sheet_data, f'{{{NS}}}row', {'r': '9'})
        return ET.tostring(root)
    rewrite_package(path, transform)


class EmptySheetInputTest(unittest.TestCase):
    def test_empty_sheets_with_placeholder_or_missing_dimensions(self):
        for dimension in ('A1:A1', None, 'A1:B2', 'A1:XFD100'):
            for empty_rows in (False, True):
                with self.subTest(dimension=dimension, empty_rows=empty_rows), TemporaryDirectory() as tmp:
                    path = Path(tmp) / 'empty.xlsx'
                    workbook = Workbook()
                    workbook.save(path)
                    workbook.close()
                    alter_dimensions(path, dimension, empty_rows)
                    before = path.read_bytes()
                    workbook = _open_workbook(path)
                    try:
                        self.assertEqual((1, 1), (workbook.active.max_row, workbook.active.max_column))
                        self.assertTrue(all(value is None for row in workbook.active.values for value in row))
                        self.assertIsNone(workbook.active['A1'].value)
                    finally:
                        workbook.close()
                    self.assertEqual(before, path.read_bytes())

    def test_single_cell_does_not_disappear_during_recovery(self):
        for value in ('合同号', 0, False):
            with self.subTest(value=value), TemporaryDirectory() as tmp:
                path = Path(tmp) / 'single.xlsx'
                workbook = Workbook()
                workbook.active['A1'] = value
                workbook.save(path)
                workbook.close()
                workbook = _open_workbook(path)
                try:
                    self.assertEqual(value, workbook.active['A1'].value)
                    self.assertIs(type(value), type(workbook.active['A1'].value))
                    self.assertEqual((1, 1), (workbook.active.max_row, workbook.active.max_column))
                finally:
                    workbook.close()

    def test_sparse_and_styled_cells_keep_coordinates_formats_and_types(self):
        for dimension in ('A1:A1', None, 'A1:B2', 'A1:XFD100'):
            for style_only in (False, True):
                with self.subTest(dimension=dimension, style_only=style_only), TemporaryDirectory() as tmp:
                    path = Path(tmp) / 'sparse.xlsx'
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet['F9'].font = Font(bold=True)
                    if not style_only:
                        sheet['B3'] = 12
                        sheet['B3'].number_format = '00000'
                        sheet['D7'] = datetime(2026, 9, 7)
                        sheet['D7'].number_format = 'yyyy-mm-dd'
                        sheet['E8'] = False
                        sheet['F9'] = 0
                    workbook.save(path)
                    workbook.close()
                    alter_dimensions(path, dimension)
                    workbook = _open_workbook(path)
                    try:
                        sheet = workbook.active
                        self.assertEqual((9, 6), (sheet.max_row, sheet.max_column))
                        self.assertTrue(sheet['F9'].font.bold)
                        if style_only:
                            self.assertIsNone(sheet['F9'].value)
                        else:
                            self.assertEqual(12, sheet['B3'].value)
                            self.assertEqual('00000', sheet['B3'].number_format)
                            self.assertEqual(datetime(2026, 9, 7), sheet['D7'].value)
                            self.assertEqual('yyyy-mm-dd', sheet['D7'].number_format)
                            self.assertIs(False, sheet['E8'].value)
                            self.assertEqual(0, sheet['F9'].value)
                    finally:
                        workbook.close()

    def test_sources_with_visible_and_hidden_empty_sheets_still_import(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            sources = _write_sources(root, 'source', variant='first')
            for path in sources:
                add_empty_sheets(path)
                alter_dimensions(path, 'A1:B2')
            originals = {path: path.read_bytes() for path in sources}
            result = _run(sources, root / 'result.xlsx')
            self.assertEqual(7, result.base_count)
            for path, before in originals.items():
                self.assertEqual(before, path.read_bytes())

    def test_previous_empty_sheets_with_user_and_grouped_pivots_keep_manual_inputs(self):
        for metadata in (True, False):
            with self.subTest(metadata=metadata), TemporaryDirectory() as tmp:
                root = Path(tmp)
                sources = _write_sources(root, 'source', variant='first')
                previous, output = root / 'previous.xlsx', root / 'result.xlsx'
                _run(sources, previous)
                _set_manual_inputs(previous, 'C001', 'SC-A', segment_flag=False, rpd='2026-9', cpd='2026-10', amount=0)
                add_empty_sheets(previous)
                add_user_pivots(previous, metadata)
                add_discrete_group(previous)
                alter_dimensions(previous, 'A1:B2')
                before = previous.read_bytes()
                result = _run(sources, output, previous=previous)
                self.assertEqual(7, result.base_count)
                self.assertEqual(before, previous.read_bytes())
                workbook = load_workbook(output, data_only=True)
                try:
                    row = _base_rows(workbook['基表'])[('C001', 'SC-A')]
                    self.assertIs(False, row['最终收入分段类别'])
                    self.assertEqual(0, row['调整金额'])
                    self.assertEqual(0, row['最终收入预测'])
                    self.assertEqual('2026-09', row['最终收入年月（按RPD）'])
                    self.assertEqual('2026-10', row['最终收入年月（按CPD）'])
                    self.assertEqual({'RegionalRevenueRPD', 'RegionalRevenueCPD'}, {p.name for s in workbook for p in s._pivots})
                finally:
                    workbook.close()

    def test_fully_empty_previous_reports_missing_header_not_python_exception(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'empty.xlsx'
            workbook = Workbook()
            workbook.active.title = '基表'
            workbook.save(path)
            workbook.close()
            issues = IssueLog()
            previous = ExcelInputAdapter().read_previous(path, load_config(CONFIG), issues)
            self.assertFalse(previous.usable)
            self.assertFalse(previous.rows)
            self.assertIn('PREVIOUS_HEADER_NOT_FOUND', {issue.code for issue in issues.items})

    def test_fully_empty_source_reports_missing_business_sheets(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'empty.xlsx'
            workbook = Workbook()
            workbook.save(path)
            workbook.close()
            issues = IssueLog()
            source = ExcelInputAdapter().read_source(
                SourceFiles(path, path, path, path), load_config(CONFIG), issues,
            )
            self.assertTrue(all(not rows for rows in source.rows.values()))
            self.assertEqual(4, sum(issue.code == 'SHEET_ROLE_NOT_FOUND' for issue in issues.items))

    def test_malformed_worksheet_is_not_silently_treated_as_empty(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'malformed.xlsx'
            workbook = Workbook()
            workbook.save(path)
            workbook.close()
            rewrite_package(path, lambda name, data: data.replace(b'</sheetData>', b'') if name == 'xl/worksheets/sheet1.xml' else data)
            with self.assertRaisesRegex(WorkbookReadError, '工作簿无法读取'):
                _open_workbook(path)
