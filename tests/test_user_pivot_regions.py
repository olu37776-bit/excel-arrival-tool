from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from revenue_tool.adapters.excel_reader import ExcelInputAdapter, _open_workbook
from revenue_tool.adapters.pivot_input_view import business_sheet
from revenue_tool.adapters.pivot_smoke_fixture import add_pivot_fixture
from revenue_tool.config import load_config
from revenue_tool.domain.models import IssueLog, WorkbookReadError
from tests.test_grouped_pivot_input import add_discrete_group, rewrite_package
from tests.test_manual_revenue_forecast import _set_manual_inputs
from tests.test_pipeline import CONFIG, _base_rows, _run, _write_sources


def add_user_pivots(path, metadata=True):
    wb = load_workbook(path)
    try:
        base = wb['基表']
        headers = [c.value for c in base[1]]
        fake = [c.value for c in base[2]]
        fake[headers.index('合同号')] = 'PIVOT-ONLY'
        template = add_pivot_fixture(
            wb, base, location_ref='AS3:CF5', name='UserRight'
        )
        base['AS1'], base['AT1'] = '透视筛选', 'USER'
        for offset, value in enumerate(headers):
            base.cell(3, 45 + offset, value)
            base.cell(4, 45 + offset, fake[offset])
        lower = deepcopy(template)
        lower.name = 'UserBelow'
        lower.location.ref = 'A12:AN14'
        lower.location.firstHeaderRow = 0
        lower.location.firstDataRow = 1
        lower.location.rowPageCount = 0
        lower.location.colPageCount = 0
        lower.pageFields = []
        base.add_pivot(lower)
        for offset, value in enumerate(headers):
            base.cell(12, 1 + offset, value)
            base.cell(13, 1 + offset, fake[offset])
        # Real business row after the lower pivot, including a manual 0.
        real = [c.value for c in base[2]]
        real[headers.index('合同号')] = 'REAL-AFTER-PIVOT'
        real[headers.index('调整金额')] = 0
        for column, value in enumerate(real, 1):
            base.cell(18, column, value)
        if metadata:
            meta = wb['_tool_meta']
            meta.append(['REAL-AFTER-PIVOT', real[headers.index('履行供应中心')], 'DEMAND_CENTER'])
        else:
            del wb['_tool_meta']
        wb.save(path)
    finally:
        wb.close()


class UserPivotRegionsTest(unittest.TestCase):
    def test_same_sheet_right_and_below_pivots_preserve_real_rows_and_inputs(self):
        for metadata in (True, False):
            with self.subTest(metadata=metadata), TemporaryDirectory() as tmp:
                root = Path(tmp)
                sources = _write_sources(root, 'src', variant='first')
                old, new = root / 'old.xlsx', root / 'new.xlsx'
                _run(sources, old)
                _set_manual_inputs(old, 'C001', 'SC-A', segment_flag=False, rpd='2026-9', cpd='2026-10', amount=0)
                add_user_pivots(old, metadata)
                add_discrete_group(old)
                before = old.read_bytes()
                issues = IssueLog()
                previous = ExcelInputAdapter().read_previous(old, load_config(CONFIG), issues)
                self.assertTrue(previous.usable, issues)
                contracts = {r.values['contract_no'] for r in previous.rows.values()}
                self.assertNotIn('PIVOT-ONLY', contracts)
                self.assertIn('REAL-AFTER-PIVOT', contracts)
                inherited = next(r.values for r in previous.rows.values() if r.values['contract_no'] == 'C001' and r.values['supply_center'] == 'SC-A')
                self.assertIs(False, inherited['manual_revenue_segment'])
                self.assertEqual(0, inherited['manual_revenue_month'])
                _run(sources, new, previous=old)
                wb = load_workbook(new, data_only=True)
                try:
                    actual = _base_rows(wb['基表'])[('C001', 'SC-A')]
                    self.assertEqual('2026-09', actual['最终收入年月（按RPD）'])
                    self.assertEqual('2026-10', actual['最终收入年月（按CPD）'])
                    self.assertEqual(0, actual['最终收入预测'])
                    self.assertFalse(any(s._pivots for s in wb))
                finally:
                    wb.close()
                self.assertEqual(before, old.read_bytes())

    def test_filter_cells_and_neighbors_are_distinguished(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); path = root / 'old.xlsx'
            _run(_write_sources(root, 'src', variant='first'), path)
            wb = load_workbook(path)
            sheet = wb['基表']
            add_pivot_fixture(wb, sheet, location_ref='AS3:AX8', name='UserWithFilter')
            sheet['AS1'] = '合同号';sheet['AT1'] = 'C001'
            sheet['AR1'] = '左侧真实备注';sheet['AU1'] = '右侧真实备注';sheet['AS2'] = '间隔行真实备注'
            wb.save(path);wb.close()
            wb = _open_workbook(path)
            try:
                view = business_sheet(wb, wb['基表'])
                self.assertIsNone(view['AS1'].value)
                self.assertIsNone(view['AT1'].value)
                self.assertEqual('左侧真实备注', view['AR1'].value)
                self.assertEqual('右侧真实备注', view['AU1'].value)
                self.assertEqual('间隔行真实备注', view['AS2'].value)
            finally:
                wb.close()

    def test_bad_pivot_location_stops_instead_of_silently_inheriting_results(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp);path = root / 'bad.xlsx'
            _run(_write_sources(root, 'src', variant='first'), path)
            add_user_pivots(path)
            def break_layout(name, data):
                if name.startswith('xl/pivotTables/') and name.endswith('.xml') and b'UserRight' in data:
                    return data.replace(b'AS3:CF5', b'not-a-range')
                return data
            rewrite_package(path, break_layout)
            with self.assertRaisesRegex(WorkbookReadError, '透视区域无法识别'):
                ExcelInputAdapter().read_previous(path, load_config(CONFIG), IssueLog())
