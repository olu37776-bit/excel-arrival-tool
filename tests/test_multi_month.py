"""All-month horizontal report, replacing the former selected-month contract."""
from copy import deepcopy
from decimal import Decimal
import json
import os
from pathlib import Path
import shutil
import subprocess
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from revenue_tool.adapters.excel_writer import ExcelOutputAdapter
from revenue_tool.adapters.regional_pivot import SUMMARY_SHEETS
from revenue_tool.config import load_config
from revenue_tool.domain.models import BaseRow, IssueLog
from revenue_tool.services.final_revenue import calculate_final_values
from revenue_tool.services.regional_summary import LABELS, EXCLUDED, build_regional_summary, discover_months
from tests.test_pipeline import CONFIG
from tests.test_regional_summary import example_rows


class HorizontalMonthsTest(unittest.TestCase):
    def test_all_actual_months_include_years_and_manual_overrides(self):
        rows = example_rows()
        self.assertEqual(('2025-09', '2026-01', '2026-08', '2026-09', '2026-10'), discover_months(rows))
        rows[0].values['manual_revenue_forecast_rpd'] = '2027-02'
        self.assertIn('2027-02', discover_months(rows))
        self.assertEqual((), discover_months([BaseRow({'revenue_month_rpd': '0000-01'})]))
        many = [BaseRow({'revenue_month_rpd': f'{year}-{m:02d}'}) for year in (2025, 2026) for m in range(1, 13)]
        self.assertEqual(24, len(discover_months(many)))

    def test_monthly_totals_and_every_cell_detail(self):
        rows = example_rows()
        for mode, expected in [('rpd', [20, -5, 0, 7, 22]), ('cpd', [100, -5, 0, 7, 102])]:
            summary = build_regional_summary(rows, mode)
            self.assertEqual(expected, summary.totals['A']['2026-09'])
            self.assertEqual([900, 0, 0, 0, 900], summary.totals['A']['2025-09'])
            self.assertEqual([800, 0, 0, 0, 800], summary.totals['A']['2026-10'])
            for region in [*summary.regions, None]:
                for month in summary.months:
                    for i, label in enumerate(LABELS):
                        entries = [e for e in summary.entries if e.month == month and e.bucket == label and (region is None or e.region == region)]
                        expected = summary.totals[region][month][i] if region else sum(t[month][i] for t in summary.totals.values())
                        self.assertEqual(expected, sum((e.amount for e in entries), Decimal(0)))
                        self.assertEqual(len(entries), len({e.base_row for e in entries}))

    def test_native_headers_filters_caches_and_editable_base(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'horizontal.xlsx'
            rows = example_rows()
            ExcelOutputAdapter().write(path, rows, [], [], [], IssueLog(), load_config(CONFIG))
            wb = load_workbook(path)
            try:
                self.assertFalse(wb['基表'].protection.sheet)
                self.assertIsNotNone(wb['基表'].auto_filter.ref)
                self.assertEqual(45, wb['_summary_source'].max_column)
                self.assertEqual(4 * len(rows) + 5 * 5 + 1, wb['_summary_source'].max_row)
                self.assertEqual(set(SUMMARY_SHEETS.values()), {s.title for s in wb if s._pivots})
                for mode_index, (mode, title) in enumerate(SUMMARY_SHEETS.items()):
                    sheet = wb[title]
                    self.assertFalse(sheet.protection.sheet)
                    self.assertFalse(sheet.row_dimensions[8].hidden)
                    self.assertFalse(sheet.row_dimensions[4].hidden)
                    self.assertEqual(26, sheet.max_column)
                    self.assertEqual('小计', sheet['A12'].value)
                    pivot = sheet._pivots[0]
                    self.assertEqual(2, len(pivot.colFields))
                    self.assertTrue(pivot.showHeaders)
                    self.assertTrue(pivot.enableDrill)
                    self.assertFalse(pivot.colGrandTotals)
                    summary = build_regional_summary(rows, mode)
                    for mi, month in enumerate(summary.months):
                        start = 2 + mi * 5
                        self.assertEqual(month, sheet.cell(7, start).value)
                        self.assertEqual(list(LABELS), [sheet.cell(8, start+j).value for j in range(5)])
                        for ri, region in enumerate(summary.regions):
                            for bi in range(5):
                                detail = [r._fields for r in pivot.cache.records.r if r._fields[40].v == ri
                                    and r._fields[41].v == mi and r._fields[42].v == bi and r._fields[44].v == mode_index]
                                self.assertEqual(sheet.cell(9+ri, start+bi).value, sum(r[38].v for r in detail))
                                self.assertEqual(len(detail), len({r[43].v for r in detail}))
            finally:
                wb.close()
            with ZipFile(path) as archive:
                ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                base = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
                self.assertIsNone(base.find('x:sheetProtection', ns))
                self.assertNotIn('前期累计', archive.read('xl/pivotCache/pivotCacheDefinition1.xml').decode())

    def test_no_valid_months_do_not_invent_a_month(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'empty.xlsx'
            ExcelOutputAdapter().write(path, [], [], [], [], IssueLog(), load_config(CONFIG))
            wb = load_workbook(path)
            try:
                for title in SUMMARY_SHEETS.values():
                    self.assertIn('暂无有效收入年月', wb[title]['A2'].value)
                    self.assertFalse(wb[title]._pivots)
            finally:
                wb.close()

    def test_live_edit_without_f9_and_native_monthly_details(self):
        office = shutil.which('libreoffice') or shutil.which('soffice')
        uno_python = Path('/usr/bin/python3')
        if not office or not uno_python.exists() or subprocess.run([str(uno_python), '-c', 'import uno'], capture_output=True).returncode:
            if os.environ.get('REQUIRE_FORMULA_ENGINE') == '1':
                self.fail('Native automatic calculation engine required')
            self.skipTest('Native calculation verification runs in Linux job')
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / 'horizontal.xlsx'
            rows = example_rows()
            config = load_config(CONFIG)
            ExcelOutputAdapter().write(path, rows, [], [], [], IssueLog(), config)
            ids = [c['id'] for c in config.base_columns]
            def edit(field, value):
                return ['基表', f'{get_column_letter(ids.index(field)+1)}2', value]
            manual = ['manual_revenue_forecast_rpd', 'manual_revenue_forecast_cpd', 'manual_revenue_segment', 'manual_revenue_month']
            plan = [[], [edit(manual[0], '9月'), edit(manual[1], '10月'), edit(manual[2], '特殊处理'), edit(manual[3], -10)],
                    [edit(manual[2], False), edit(manual[3], 0)], [edit(f, None) for f in manual], [edit(manual[0], '2027-03')]]
            (root / 'plan.json').write_text(json.dumps(plan), encoding='utf-8')
            result = subprocess.run([str(uno_python), str(Path(__file__).parent / 'support/verify_native_pivot.py'),
                office, str(path), str(root / 'plan.json'), str(root / 'result.json')], capture_output=True, text=True, timeout=150)
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            snapshots = json.loads((root / 'result.json').read_text())
            initial_months = discover_months(rows)
            for stage, snapshot in enumerate(snapshots):
                current = deepcopy(rows)
                if stage in (1, 2):
                    current[0].values.update(zip(manual, ['9月', '10月', '特殊处理' if stage == 1 else False, -10 if stage == 1 else 0]))
                if stage == 4:
                    current[0].values[manual[0]] = '2027-03'
                expected_final = calculate_final_values(current[0].values)
                for field, value in expected_final.items():
                    # UNO exposes boolean result cells as 0/1 in getDataArray.
                    self.assertEqual(value, snapshot['base'][1][ids.index(field)], (stage, field))
                for mode, title in SUMMARY_SHEETS.items():
                    actual = snapshot['summaries'][title]
                    wanted_months = tuple(sorted(set(initial_months) | set(discover_months(current))))
                    self.assertEqual(list(wanted_months), actual['months'])
                    summary = build_regional_summary(current, mode, wanted_months)
                    for region in [*summary.regions, '小计']:
                        for month in wanted_months:
                            for label in LABELS:
                                entries = [e for e in summary.entries if e.month == month and e.bucket == label and (region == '小计' or e.region == region)]
                                cell = actual['cells'][region][month][label]
                                self.assertEqual(sorted(current[e.base_row-2].values['contract_no'] for e in entries), sorted(cell['contracts']), (stage, mode, month, region, label))
                                self.assertEqual(float(sum((e.amount for e in entries), Decimal(0))), cell['value'])
