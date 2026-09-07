"""Native OOXML pivots with saved caches and formula-backed drill-through sources.

One pivot column = one bucket. Current rows also have a subtotal bucket copy;
row grand totals are disabled, so they cannot accidentally sum both copies.
Never use six conditional measures: their Show Details would include unrelated
months/segments from the entire region. Source sheets are not business facts.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.fields import Boolean, Index, Missing, Number, Text
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import (
    DataField, FieldItem, Location, PivotField, RowColField, RowColItem,
    TableDefinition, PivotTableStyle, PageField,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.utils.datetime import to_excel

from revenue_tool.services.regional_summary import (
    EMPTY_REGION, EXCLUDED, SEGMENTS, build_regional_summary,
)

SUMMARY_SHEETS = {"rpd": "RPD地区收入汇总", "cpd": "CPD地区收入汇总"}
SOURCE_SHEETS = {"rpd": "_summary_source", "cpd": "_summary_source"}


def _q(value):
    return '"' + str(value).replace('"', '""') + '"'


def _cache_value(value):
    if value is None or value == "":
        return Missing()
    if isinstance(value, bool):
        return Boolean(v=value)
    if isinstance(value, (date, datetime)):
        return Number(v=to_excel(value))
    if isinstance(value, (int, float, Decimal)):
        return Number(v=float(value))
    return Text(v=str(value))


def _bucket_formula(refs, summary, subtotal):
    month, amount, segment = (refs[f"final_revenue_month_{summary.mode}"],
                              refs["final_revenue_forecast"], refs["final_revenue_segment"])
    labels = summary.labels
    other = _q(labels[4])
    for position in range(2, -1, -1):
        other = f'IF({segment}={_q(SEGMENTS[position])},{_q(labels[position + 1])},{other})'
    valid = (f'AND(ISNUMBER({amount}),ISTEXT({month}),LEN({month})=7,'
             f'LEFT({month},5)={_q(summary.month[:4] + "-")},'
             f'IFERROR(TEXT(VALUE(RIGHT({month},2)),"00")=RIGHT({month},2),FALSE),'
             f'RIGHT({month},2)>="01",RIGHT({month},2)<="12")')
    current = _q(labels[5]) if subtotal else other
    earlier = _q(EXCLUDED) if subtotal else _q(labels[0])
    return (f'=IFERROR(IF({valid},IF({month}={_q(summary.month)},{current},'
            f'IF({month}<{_q(summary.month)},{earlier},{_q(EXCLUDED)})),'
            f'{_q(EXCLUDED)}),{_q(EXCLUDED)})')


def write_regional_pivots(workbook, rows, config, month):
    base = workbook[config.output["sheets"]["base"]]
    columns = config.base_columns
    ids = [c["id"] for c in columns]
    headers = [c["name"] for c in columns] + ["汇总地区部", "汇总项目", "基表行号", "收入口径"]
    count = len(columns)
    if 4 * len(rows) + 1 > 1048576:
        raise ValueError("基表超过262143行，双口径明细源超出Excel工作表行数限制")
    report_sheets = {mode: workbook.create_sheet(title) for mode, title in SUMMARY_SHEETS.items()}
    source = workbook.create_sheet("_summary_source")
    source.append(headers)
    source.sheet_state = "hidden"
    source.freeze_panes = "A2"
    first = build_regional_summary(rows, month, "rpd")
    regions = first.regions
    buckets = [*first.labels, EXCLUDED]
    region_ids = {v: i for i, v in enumerate(regions)}
    bucket_ids = {v: i for i, v in enumerate(buckets)}
    fields = [CacheField(name=name, sharedItems=SharedItems()) for name in headers]
    fields[count].sharedItems = SharedItems(_fields=[Text(v=x) for x in regions], containsString=True)
    fields[count + 1].sharedItems = SharedItems(_fields=[Text(v=x) for x in buckets], containsString=True)
    fields[count + 3].sharedItems = SharedItems(_fields=[Text(v=x) for x in ("RPD", "CPD")], containsString=True)
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(
            ref=f"A1:{get_column_letter(len(headers))}{max(1, 4 * len(rows) + 1)}", sheet=source.title)),
        cacheFields=fields, recordCount=4 * len(rows), saveData=True,
        refreshOnLoad=False, enableRefresh=True, backgroundQuery=False,
        createdVersion=6, refreshedVersion=6, minRefreshableVersion=3,
    )
    cache.records = RecordList()
    records = cache.records.r
    for mode in ("rpd", "cpd"):
        summary = build_regional_summary(rows, month, mode)
        sheet = report_sheets[mode]
        mode_index = 0 if mode == "rpd" else 1
        for source_number, entry in enumerate(summary.entries, len(records) + 2):
            values = rows[entry.base_row - 2].values
            for index, field in enumerate(ids, 1):
                ref = f'{quote_sheetname(base.title)}!{get_column_letter(index)}{entry.base_row}'
                cell = source.cell(source_number, index, f'=IF(ISBLANK({ref}),"",{ref})')
                cell.number_format = base.cell(entry.base_row, index).number_format
            refs = {field: f"{get_column_letter(i)}{source_number}" for i, field in enumerate(ids, 1)}
            region = refs["region"]
            source.cell(source_number, count + 1,
                        f'=IF(LEN(TRIM({region}&""))=0,{_q(EMPTY_REGION)},{region}&"")')
            source.cell(source_number, count + 2, _bucket_formula(refs, summary, entry.subtotal_copy))
            source.cell(source_number, count + 3, entry.base_row)
            source.cell(source_number, count + 4, mode.upper())
            records.append(Record(_fields=[*(_cache_value(values.get(f)) for f in ids),
                Index(v=region_ids[entry.region]), Index(v=bucket_ids[entry.bucket]),
                Number(v=entry.base_row), Index(v=mode_index)]))
        source.auto_filter.ref = source.dimensions
        pivot_fields = [PivotField(defaultSubtotal=False) for _ in headers]
        pivot_fields[count] = PivotField(axis="axisRow", defaultSubtotal=False, showAll=True,
            name="地区部", items=[FieldItem(x=i) for i in range(len(regions))] + [FieldItem(t="grand")])
        pivot_fields[count + 1] = PivotField(axis="axisCol", defaultSubtotal=False,
            showAll=True, sortType="manual", items=[FieldItem(x=i, h=(i == 6)) for i in range(7)])
        pivot_fields[count + 3] = PivotField(axis="axisPage", defaultSubtotal=False,
            items=[FieldItem(x=0), FieldItem(x=1)])
        amount_index = ids.index("final_revenue_forecast")
        pivot_fields[amount_index].dataField = True
        pivot = TableDefinition(
            name=f"RegionalRevenue{mode.upper()}", cacheId=1, dataCaption="最终收入预测",
            grandTotalCaption="小计", rowHeaderCaption="地区部", colHeaderCaption="汇总项目",
            location=Location(ref=f"A6:G{8 + len(regions)}", firstHeaderRow=1, firstDataRow=2, firstDataCol=1, rowPageCount=1, colPageCount=1),
            pivotFields=pivot_fields, rowFields=[RowColField(x=count)], colFields=[RowColField(x=count + 1)],
            pageFields=[PageField(fld=count + 3, item=mode_index)],
            rowItems=[RowColItem(x=[Index(v=i)]) for i in range(len(regions))] + [RowColItem(t="grand", x=[Index(v=0)])],
            colItems=[RowColItem(x=[Index(v=i)]) for i in range(6)],
            dataFields=[DataField(name="最终收入预测汇总", fld=amount_index, subtotal="sum", numFmtId=4)],
            rowGrandTotals=False, colGrandTotals=True, enableDrill=True,
            subtotalHiddenItems=False, showHeaders=True, compact=False, compactData=False,
            outline=False, gridDropZones=False, showDrill=False, showEmptyRow=True,
            showEmptyCol=True, missingCaption="0.00", showMissing=True,
            createdVersion=6, updatedVersion=6, minRefreshableVersion=3,
            pivotTableStyleInfo=PivotTableStyle(name="PivotStyleMedium9", showRowHeaders=True,
                showColHeaders=True, showRowStripes=True, showColStripes=False, showLastColumn=True),
        )
        pivot.cache = cache
        sheet.add_pivot(pivot)
        sheet.merge_cells("A1:G1")
        sheet["A1"] = f"{month} {mode.upper()}地区收入汇总"
        sheet["A1"].font = Font(size=16, bold=True, color="1F4E78")
        sheet.row_dimensions[1].height = 29
        sheet.merge_cells("A2:G2")
        sheet["A2"] = "双击金额查看明细。修改黄色人工字段后，先按Ctrl+Alt+F9重算，再点“数据→全部刷新”。"
        sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[2].height = 32
        sheet.merge_cells("A3:G3")
        sheet["A3"] = f"统计月份：{month}；仅汇总本年截至当月。初次生成未纳入：{summary.excluded_count}条（月份不在范围、空白或金额无效）；辅助源请勿直接求和。"
        sheet["A3"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet["A3"].font = Font(size=10, color="666666")
        sheet.row_dimensions[3].height = 30
        sheet["A4"] = "收入口径"
        sheet["B4"] = mode.upper()
        sheet["A6"] = "最终收入预测"
        sheet["B6"] = "汇总项目"
        for index, label in enumerate(["地区部", *summary.labels], 1):
            sheet.cell(7, index, label)
        for number, region in enumerate(regions, 8):
            sheet.cell(number, 1, region)
            for column, value in enumerate(summary.totals[region], 2):
                sheet.cell(number, column, float(value))
        last = 8 + len(regions)
        sheet.cell(last, 1, "小计")
        for column in range(6):
            sheet.cell(last, column + 2, float(sum((v[column] for v in summary.totals.values()), Decimal("0.00"))))
        for row in sheet.iter_rows(min_row=6, max_row=last, max_col=7):
            for cell in row:
                cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column == 1 else "right")
                if cell.row in (6, 7):
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.font = Font(bold=True, color="FFFFFF")
                else:
                    if cell.column > 1:
                        cell.number_format = "#,##0.00"
                    if cell.row == last or cell.column == 7:
                        cell.fill = PatternFill("solid", fgColor="D9EAF7")
                        cell.font = Font(bold=True)
                sheet.row_dimensions[cell.row].height = 22
        sheet.column_dimensions["A"].width = 26
        for column in "BCDEFG":
            sheet.column_dimensions[column].width = 19
        sheet.freeze_panes = "B8"
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:G{last}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
