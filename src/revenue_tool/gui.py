from __future__ import annotations

import argparse
import os
from pathlib import Path
import sys

from revenue_tool.config import load_config


def default_config_path() -> Path:
    """Return the bundled config path for source and frozen builds."""
    if getattr(sys, "frozen", False):
        return Path(getattr(sys, "_MEIPASS")) / "config" / "default.json"
    return Path(__file__).resolve().parents[2] / "config" / "default.json"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Excel tool graphical launcher")
    parser.add_argument(
        "--config",
        default=str(default_config_path()),
        help="Configuration JSON path (defaults to the bundled configuration)",
    )
    parser.add_argument("--smoke-test", action="store_true", help=argparse.SUPPRESS)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.smoke_test:
        import tkinter  # noqa: F401 - verifies the frozen GUI runtime
        from tempfile import TemporaryDirectory
        from openpyxl import load_workbook
        from revenue_tool.adapters.excel_writer import ExcelOutputAdapter
        from revenue_tool.adapters.pivot_smoke_fixture import add_pivot_fixture
        from revenue_tool.domain.models import BaseRow, IssueLog
        from revenue_tool.services.final_revenue import calculate_final_values

        config = load_config(args.config)
        values = dict(contract_no="SMOKE", region="测试地区", supply_center="深供",
                      revenue_month_rpd="2026-09", revenue_month_cpd="2026-09",
                      revenue_segment="订未发", revenue_forecast=1,
                      manual_revenue_segment=False, manual_revenue_month=0)
        values.update(calculate_final_values(values))
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "smoke.xlsx"
            ExcelOutputAdapter().write(path, [BaseRow(values)], [], [], [], IssueLog(), config)
            workbook = load_workbook(path)
            try:
                if any(sheet._pivots for sheet in workbook.worksheets):
                    raise RuntimeError("Windows EXE正常输出不应包含系统透视表")
                if "_summary_source" in workbook.sheetnames:
                    raise RuntimeError("Windows EXE正常输出不应包含透视辅助源")
                if workbook.calculation.forceFullCalc or workbook.calculation.fullCalcOnLoad:
                    raise RuntimeError("Windows EXE正常输出不应启用强制全量重算")
            finally:
                workbook.close()
            workbook = load_workbook(path, data_only=True)
            try:
                base = workbook[config.output["sheets"]["base"]]
                indexes = {c["id"]: i for i, c in enumerate(config.base_columns, 1)}
                for field, expected in calculate_final_values(values).items():
                    if base.cell(2, indexes[field]).value != expected:
                        raise RuntimeError("Windows EXE最终字段计算值自检失败")
            finally:
                workbook.close()
            # Input compatibility still needs a real native user PivotTable even
            # though normal output no longer generates any system pivots.
            workbook = load_workbook(path)
            try:
                base = workbook[config.output['sheets']['base']]
                add_pivot_fixture(
                    workbook,
                    base,
                    location_ref='AS1:AT3',
                    name='UserPivotSmoke',
                )
                base['AS1'], base['AT1'] = '合同号', '金额'
                base['AS2'], base['AT2'] = 'PIVOT-ONLY', 999
                # Empty user sheets must not break dimension recovery/import.
                for state in ('visible', 'hidden', 'veryHidden'):
                    workbook.create_sheet('Empty-' + state).sheet_state = state
                workbook.save(path)
            finally:
                workbook.close()
            # The packaged EXE must read grouped legacy pivot caches too.
            # This temporary fixture reproduces openpyxl 3.1.5's Nested error.
            from zipfile import ZipFile, ZIP_DEFLATED
            from xml.etree import ElementTree as ET
            from revenue_tool.adapters.excel_reader import ExcelInputAdapter
            namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
            with ZipFile(path) as archive:
                parts = [(info, archive.read(info.filename)) for info in archive.infolist()]
            with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
                for info, data in parts:
                    if info.filename.startswith('xl/pivotCache/pivotCacheDefinition') and info.filename.endswith('.xml'):
                        tree = ET.fromstring(data)
                        field = tree.find(f'{{{namespace}}}cacheFields')[0]
                        group = ET.SubElement(field, f'{{{namespace}}}fieldGroup', {'base': '0'})
                        discrete = ET.SubElement(group, f'{{{namespace}}}discretePr', {'count': '1'})
                        ET.SubElement(discrete, f'{{{namespace}}}x', {'v': '0'})
                        data = ET.tostring(tree)
                    archive.writestr(info, data)
            previous = ExcelInputAdapter().read_previous(path, config, IssueLog())
            if not previous.usable or len(previous.rows) != 1:
                raise RuntimeError('Windows EXE分组透视上期读取自检失败')
            inherited = next(iter(previous.rows.values())).values
            if inherited['manual_revenue_segment'] is not False or inherited['manual_revenue_month'] != 0:
                raise RuntimeError('Windows EXE人工字段继承自检失败')
        if sys.platform == 'win32' or os.environ.get('DISPLAY'):
            from revenue_tool.gui_app import RevenueApp
            root = tkinter.Tk()
            try:
                app = RevenueApp(root, args.config)
                root.update()
                if hasattr(app, 'months') or hasattr(app, 'year'):
                    raise RuntimeError('GUI不应包含月份选择')
            finally:
                root.destroy()
        return 0

    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except ImportError:
        print(
            "Tkinter is unavailable. Reinstall Python from python.org with "
            "the Tcl/Tk option enabled.",
            file=sys.stderr,
        )
        return 2

    from revenue_tool.gui_app import RevenueApp
    root = tk.Tk()
    RevenueApp(root, args.config)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
